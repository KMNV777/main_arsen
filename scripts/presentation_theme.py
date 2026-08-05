"""Деловой стиль слайдов для презентаций (python-pptx)."""

from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Палитра
NAVY = RGBColor(0x00, 0x33, 0x66)
ACCENT = RGBColor(0x00, 0x70, 0xC0)
BG_LIGHT = RGBColor(0xF4, 0xF6, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x33, 0x33, 0x33)
TEXT_MUTED = RGBColor(0x66, 0x66, 0x66)
GREEN = RGBColor(0x1B, 0x7A, 0x3A)
RED = RGBColor(0xC0, 0x00, 0x00)
ORANGE = RGBColor(0xE6, 0x7E, 0x22)
TABLE_HEADER = NAVY
TABLE_ALT = RGBColor(0xE8, 0xEE, 0xF4)

FONT = "Calibri"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
HEADER_H = Inches(0.72)
FOOTER_H = Inches(0.38)
CONTENT_TOP = Inches(0.95)
FOOTER_Y = Inches(7.05)


def set_font(run, size: int = 14, bold: bool = False, color: RGBColor | None = None, name: str = FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _send_to_back(slide, shape):
    sp_tree = slide.shapes._spTree
    sp = shape._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)


def apply_frame(slide, title: str = "", page: int | None = None, *, dark: bool = False):
    """Шапка, фон и подвал в деловом стиле."""
    bg_color = NAVY if dark else BG_LIGHT
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = bg_color
    bg.line.fill.background()
    _send_to_back(slide, bg)

    if not dark:
        header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, HEADER_H)
        header.fill.solid()
        header.fill.fore_color.rgb = NAVY
        header.line.fill.background()

        if title:
            tb = slide.shapes.add_textbox(Inches(0.55), Inches(0.14), Inches(11.5), Inches(0.5))
            p = tb.text_frame.paragraphs[0]
            r = p.add_run()
            r.text = title
            set_font(r, 22, True, WHITE)

        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, HEADER_H, SLIDE_W, Inches(0.04))
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT
        accent.line.fill.background()

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.5), FOOTER_Y, Inches(12.33), Inches(0.01))
        line.fill.solid()
        line.fill.fore_color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        line.line.fill.background()

        foot = slide.shapes.add_textbox(Inches(0.55), Inches(7.12), Inches(8), Inches(0.28))
        fr = foot.text_frame.paragraphs[0].add_run()
        fr.text = "ЦМ ИМА × ЕАРМ  |  Сравнительный анализ вариантов интеграции"
        set_font(fr, 8, color=TEXT_MUTED)

        if page is not None:
            pn = slide.shapes.add_textbox(Inches(12.2), Inches(7.1), Inches(0.8), Inches(0.3))
            pp = pn.text_frame.paragraphs[0]
            pp.alignment = PP_ALIGN.RIGHT
            pr = pp.add_run()
            pr.text = str(page)
            set_font(pr, 9, color=TEXT_MUTED)
    elif title:
        tb = slide.shapes.add_textbox(Inches(0.6), Inches(2.8), Inches(12), Inches(1.2))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        set_font(r, 30, True, WHITE)

    return slide


def style_table_header(cell, text: str):
    cell.text = text
    cell.fill.solid()
    cell.fill.fore_color.rgb = TABLE_HEADER
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for p in cell.text_frame.paragraphs:
        p.alignment = PP_ALIGN.CENTER
        for r in p.runs:
            set_font(r, 9, True, WHITE)


def style_table_cell(cell, text: str, *, bold: bool = False, alt: bool = False):
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    if alt:
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_ALT
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            set_font(r, 8, bold, TEXT_DARK)
